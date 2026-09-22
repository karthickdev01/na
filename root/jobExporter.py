from datetime import datetime
from pathlib import Path
from urllib.parse import urljoin

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from config import G_EXPORTED_DATA_PATH, G_NAKURI_WEB_URL

COLUMNS = [
    ("S.No.", None),
    ("Job ID", "jobId"),
    ("Job Title", "title"),
    ("Company", "companyName"),
    ("Experience", "experienceText"),
    ("Location", "location"),
    ("Salary", "salary"),
    ("Skills", "tagsAndSkills"),
    ("ATS Score", "atsScore"),
    ("ATS Reason", "atsReason"),
    ("Job Description", "jobDescription"),
    ("Apply Today", "todaysJob"),
    ("Walk-in", "walkinJob"),
    ("Questionnaire", "questionnaireIdPresent"),
    ("Company Apply", "companyApplyJob"),
    ("Company Apply URL", "companyApplyUrl"),
    ("Apply Redirect URL", "applyRedirectUrl"),
    ("Link", "jdURL"),
]

URL_FIELDS = {"jdURL", "companyApplyUrl", "applyRedirectUrl"}

CELL_BORDER = Border(
    left=Side(style="thin", color="D9E2F3"),
    right=Side(style="thin", color="D9E2F3"),
    top=Side(style="thin", color="D9E2F3"),
    bottom=Side(style="thin", color="D9E2F3"),
)


def _placeholder_value(job: dict, placeholder_type: str) -> str:
    for placeholder in job.get("placeholders") or []:
        if placeholder.get("type") == placeholder_type:
            return placeholder.get("label", "")
    return ""


def _salary_value(job: dict) -> str:
    salary = job.get("salaryDetail") or {}
    minimum = salary.get("minimumSalary")
    maximum = salary.get("maximumSalary")
    currency = salary.get("currency", "")
    if minimum is None and maximum is None:
        return _placeholder_value(job, "salary")
    if minimum is not None and maximum is not None:
        return f"{currency} {minimum:,} - {maximum:,}".strip()
    value = minimum if minimum is not None else maximum
    return f"{currency} {value:,}".strip()


def _export_value(job: dict, field: str):
    if field == "location":
        return _placeholder_value(job, "location")
    if field == "salary":
        return _salary_value(job)
    if field == "jobDescription":
        from root.resumeMatcher import clean_job_description

        return clean_job_description(job.get(field) or "")
    if field in URL_FIELDS:
        value = job.get(field, "")
        return urljoin(f"{G_NAKURI_WEB_URL}/", str(value)) if value else ""
    return job.get(field, "")


def exportJobs(jobs: list, filename: str = None) -> str:
    if filename is None:
        export_directory = Path(G_EXPORTED_DATA_PATH)
        export_directory.mkdir(parents=True, exist_ok=True)
        filename = export_directory / f"{datetime.now():%Y-%m-%d %H:%M:%S} - Job List.xlsx"
    else:
        filename = Path(filename)

    wb = Workbook()
    ws = wb.active
    ws.title = "Jobs"
    ws.append([header for header, _ in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.border = CELL_BORDER

    for serial_number, job in enumerate(jobs, start=1):
        row = []
        for _, field in COLUMNS:
            value = serial_number if field is None else _export_value(job, field)
            row.append(value)
        ws.append(row)

        for column_number, (_, field) in enumerate(COLUMNS, start=1):
            if field in URL_FIELDS:
                link_cell = ws.cell(row=ws.max_row, column=column_number)
                if link_cell.value:
                    link_cell.hyperlink = link_cell.value
                    link_cell.style = "Hyperlink"

    widths = {
        "S.No.": 8, "Job ID": 16, "Job Title": 30, "Company": 28,
        "Experience": 16, "Location": 24, "Salary": 24, "Skills": 42,
        "ATS Score": 12, "ATS Reason": 42, "Job Description": 80,
        "Apply Today": 14, "Walk-in": 12, "Questionnaire": 16,
        "Company Apply": 16, "Company Apply URL": 45,
        "Apply Redirect URL": 45, "Link": 55,
    }
    for i, (header, _) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(header, 14)

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = CELL_BORDER

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False

    wb.save(filename)
    print(f"Saved {len(jobs)} jobs to {filename}")
    return filename